from __future__ import annotations

import base64
import json
from pathlib import Path

from docx import Document
import pytest

from sensorlab.report_builder import (
    available_report_devices,
    build_results_report,
    canonical_condition,
    discover_degradation_sources,
)


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "/A8AAusB9Wl2nWQAAAAASUVORK5CYII="
)


def _png(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_PNG)


def _make_condition(root: Path, name: str, device: str, *, complete: bool = True) -> Path:
    condition = root / name
    paths = [
        f"JG_n03/sweeps/{device}_Vbg0.png",
        f"JG_n03/sweeps/{device}_Vbg4.png",
        "JG_n03/trends/average/delta_Vth.png",
        f"JG_H2/overlays/{device}_baseline.png",
        f"JG_H2/overlays/{device}_end.png",
        "JG_H2/trends/average/delta_Vth_H2.png",
        "JG_H2/trends/average/resp_at_Vth_H2_pct.png",
        "JG_H2/trends/average/delta_Vth_recovery.png",
        f"SE_n04/sensing/{device}.png",
        "SE_n04/trends/average/response_delta.png",
        "SE_n04/trends/average/t90_response_s.png",
        f"STEPS_n07/sensing/{device}.png",
        "STEPS_n07/response_vs_ppm.png",
        "STEPS_n07/sensitivity_vs_day.png",
        "STEPS_n07/sensitivity_vs_cycle.png",
    ]
    if complete:
        paths.extend([
            f"JG_n03/collapsed/{device}_Vbg0.png",
            f"JG_n03/collapsed/{device}_Vbg4.png",
        ])
    for relative in paths:
        _png(condition / relative)
    return condition


def _make_outline(path: Path) -> None:
    document = Document()
    document.add_paragraph("Results Outline")
    document.add_paragraph("2. Reliability Tests")
    document.add_paragraph("Reliability overview text")
    document.add_paragraph()
    document.add_paragraph()
    document.add_paragraph("2.1 High-Temperature Storage")
    document.add_paragraph("Hot-storage notes")
    document.add_paragraph()
    document.add_paragraph()
    document.add_paragraph("2.2 85°C/85% RH Test")
    document.add_paragraph("Humidity notes")
    document.add_paragraph()
    document.add_paragraph("2.3 Temperature Cycling")
    document.add_paragraph("Cycles notes")
    document.add_paragraph()
    document.save(path)


def _make_overview(root: Path) -> None:
    for relative in (
        "publication_plots/regular/stress_induced_deltaVTh.png",
        "publication_plots/normalized/normalized_H2_induced_deltaVTh.png",
        "publication_plots/normalized/normalized_maximum_H2_response_vs_stress_time.png",
        "publication_plots/regular/recovery_deltaVTh_vs_stress_time.png",
        "publication_plots/normalized/normalized_sensitivity_vs_stress_time.png",
        "publication_plots/normalized/normalized_H2_response_time_vs_stress_time.png",
    ):
        _png(root / relative)


def test_condition_normalization() -> None:
    assert canonical_condition("Hot storage") == "hot_storage"
    assert canonical_condition("125C Hot Storage") == "hot_storage"
    assert canonical_condition("85C_85RH") == "85c_85rh"
    assert canonical_condition("cycles") == "cycles"


def test_discover_multiple_degradation_roots_and_devices(tmp_path: Path) -> None:
    normal = tmp_path / "normal" / "degradation"
    cycle_root = tmp_path / "cycle-study" / "degradation"
    _make_condition(normal, "Hot storage", "ESN1216_D1T1")
    _make_condition(normal, "85C_85RH", "ESN1259_D1T2")
    _make_condition(cycle_root, "cycles", "ESN1270_D1T2")

    sources = discover_degradation_sources([normal, cycle_root])

    assert list(sorted(sources)) == ["85c_85rh", "cycles", "hot_storage"]
    assert available_report_devices(sources["hot_storage"]) == ["ESN1216_D1T1"]
    assert sources["cycles"].degradation_root == cycle_root.resolve()


def test_duplicate_condition_is_rejected(tmp_path: Path) -> None:
    first = tmp_path / "a" / "degradation"
    second = tmp_path / "b" / "degradation"
    _make_condition(first, "cycles", "ESN1_D1T1")
    _make_condition(second, "Temperature cycles", "ESN2_D1T1")

    with pytest.raises(ValueError, match="more than once"):
        discover_degradation_sources([first, second])


def test_build_complete_results_outline(tmp_path: Path, monkeypatch) -> None:
    outline = tmp_path / "Results outline.docx"
    output = tmp_path / "Results outline - SensorLab.docx"
    supplement = tmp_path / "sensorlab-output"
    normal = tmp_path / "normal" / "degradation"
    cycle_root = tmp_path / "cycles-study" / "degradation"
    _make_outline(outline)
    _make_overview(supplement)
    _make_condition(normal, "Hot storage", "ESN1216_D1T1")
    _make_condition(normal, "85C_85RH", "ESN1259_D1T2")
    _make_condition(cycle_root, "cycles", "ESN1270_D1T2")
    monkeypatch.setattr(
        "sensorlab.report_builder.generate_overview_plots",
        lambda _conditions, _output: (),
    )

    result = build_results_report(
        outline,
        [normal, cycle_root],
        supplement,
        output,
        {
            "hot_storage": "ESN1216_D1T1",
            "85c_85rh": "ESN1259_D1T2",
            "cycles": "ESN1270_D1T2",
        },
    )

    built = Document(output)
    assert len(built.tables) == 13
    assert len(built.inline_shapes) == 54
    assert result.images_added == 54
    assert result.missing_images == ()
    assert result.overview_generated == ()
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["conditions"]["cycles"]["representative_device"] == "ESN1270_D1T2"
    assert manifest["images_added"] == 54
    assert manifest["missing_images"] == []


def test_missing_report_plots_become_manifested_placeholders(tmp_path: Path, monkeypatch) -> None:
    outline = tmp_path / "outline.docx"
    output = tmp_path / "output.docx"
    supplement = tmp_path / "supplement"
    normal = tmp_path / "normal" / "degradation"
    cycles = tmp_path / "cycles" / "degradation"
    _make_outline(outline)
    _make_overview(supplement)
    _make_condition(normal, "Hot storage", "ESN1_D1T1", complete=False)
    _make_condition(normal, "85C_85RH", "ESN2_D1T1", complete=False)
    _make_condition(cycles, "cycles", "ESN3_D1T1", complete=False)
    monkeypatch.setattr(
        "sensorlab.report_builder.generate_overview_plots",
        lambda _conditions, _output: (),
    )

    result = build_results_report(outline, [normal, cycles], supplement, output)

    # Two collapse plots are absent for each of the three conditions.
    assert len(result.missing_images) == 6
    assert result.images_added == 48
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert len(manifest["missing_images"]) == 6
    assert "[Missing plot]" in " ".join(
        cell.text for table in Document(output).tables for row in table.rows for cell in row.cells
    )
