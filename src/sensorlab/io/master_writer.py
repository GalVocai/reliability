"""Write readable per-test-type master workbooks (one row per device)."""

from __future__ import annotations

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..util import unique_sheet_name

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_master(path: str, sheets: dict[str, pd.DataFrame]) -> str:
    """Write one workbook; `sheets` maps sheet-name -> per-device table."""
    used: set[str] = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sheet = unique_sheet_name(name, used)
            df.to_excel(writer, sheet_name=sheet, index=False)
            _style(writer.sheets[sheet], df)
    return path


def _style(ws, df: pd.DataFrame) -> None:
    # header styling
    for col_idx, _ in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    # column widths + numeric rounding for readability
    for col_idx, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(f"{v}") for v in df[col].head(50))) if len(df) else len(str(col))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 28)
        if pd.api.types.is_float_dtype(df[col]):
            for row_idx in range(2, len(df) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = "0.000"
    ws.freeze_panes = "A2"
