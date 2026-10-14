"""Export selected per-device, per-day SensorLab metrics for JMP.

The source is the active analysis master for each timepoint.  It is passed
through the same baseline join used by the degradation study; no plotted
points, rounded values, or aggregate means are used as inputs.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import math
import os
import re
import tempfile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .degradation import build_degradation, filter_devices


_QUALIFIERS = ("Vbg", "segment", "state_Vj")
_KEYS = ["ChipID", "DeviceID", "Day", "Condition"]


def _identifier(value: object) -> str:
    if pd.isna(value):
        raise ValueError("Missing ESN or TNumber in an analysis master.")
    try:
        number = float(value)
        if math.isfinite(number) and number.is_integer():
            return str(int(number))
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _qualifier_value(value: object) -> str:
    if pd.isna(value):
        return "missing"
    return _identifier(value)


@dataclass(frozen=True, order=True)
class MetricOption:
    sheet: str
    metric: str
    qualifiers: tuple[tuple[str, str], ...] = ()

    @property
    def column(self) -> str:
        parts = [self.sheet, self.metric]
        parts.extend(f"{name}{value}" for name, value in self.qualifiers)
        return "__".join(parts)

    @property
    def label(self) -> str:
        detail = ", ".join(f"{name}={value}" for name, value in self.qualifiers)
        return f"{self.sheet}  /  {self.metric}" + (f"  /  {detail}" if detail else "")


@dataclass
class ExportSource:
    tables: dict[tuple[str, str], pd.DataFrame]
    options: list[MetricOption]
    baseline_day: float


@dataclass
class ExportTables:
    all_data: pd.DataFrame
    post_baseline: pd.DataFrame
    checkpoint_summary: pd.DataFrame
    column_info: pd.DataFrame
    quality_notes: list[str]


def load_export_source(
    timepoints: list[tuple[object, str]], *, baseline_day: object,
    include_devices: set[tuple[str, str]] | None = None,
    exclude_cond_days: set[tuple[str, str]] | None = None,
) -> ExportSource:
    """Read every selected day using SensorLab's degradation calculation."""
    from .degradation import day_to_num

    if not timepoints:
        raise ValueError("Add at least one analyzed timepoint before exporting.")
    days = [day_to_num(day) for day, _ in timepoints]
    if not all(math.isfinite(day) for day in days) or len(set(days)) != len(days):
        raise ValueError("Timepoints need distinct numeric day values.")
    base = day_to_num(baseline_day)
    if not math.isfinite(base) or base not in days:
        raise ValueError("The selected baseline must have a valid day value.")
    if len(timepoints) < 2:
        raise ValueError("Select the baseline and at least one later timepoint.")

    results = build_degradation(
        timepoints, baseline_day=baseline_day,
        exclude_cond_days=exclude_cond_days,
    )
    tables: dict[tuple[str, str], pd.DataFrame] = {}
    choices: set[MetricOption] = set()
    for key, (tidy, _aggregate) in results.items():
        condition, sheet = key
        # The gas/recovery sweep masters duplicate the n03/H2 comparison and
        # are also omitted from the degradation study's exported sheets.
        if sheet in {"JG_n05", "JG_n06"}:
            continue
        tidy = filter_devices(tidy, include_devices).copy()
        if tidy.empty:
            continue
        tables[(condition, sheet)] = tidy
        qualifiers = [name for name in _QUALIFIERS if name in tidy.columns]
        for values in tidy[["metric", *qualifiers]].drop_duplicates().itertuples(index=False, name=None):
            metric = str(values[0])
            parts = tuple((name, _qualifier_value(value))
                          for name, value in zip(qualifiers, values[1:]))
            choices.add(MetricOption(sheet, metric, parts))
    if not choices:
        raise ValueError("No numeric metrics were found for the selected devices and days.")
    return ExportSource(tables, sorted(choices), base)


def _selected_rows(source: ExportSource, chosen: set[MetricOption]) -> pd.DataFrame:
    parts: list[pd.DataFrame] = []
    for (condition, sheet), tidy in source.tables.items():
        for option in chosen:
            if option.sheet != sheet:
                continue
            mask = tidy["metric"].astype(str).eq(option.metric)
            for name, value in option.qualifiers:
                mask &= tidy[name].map(_qualifier_value).eq(value)
            selected = tidy.loc[mask, ["ESN", "TNumber", "day_num", "value",
                                       "baseline_value"]].copy()
            if selected.empty:
                continue
            selected["ChipID"] = "ESN" + selected["ESN"].map(_identifier)
            selected["DeviceID"] = (selected["ChipID"] + "_T"
                                    + selected["TNumber"].map(_identifier))
            selected["Day"] = pd.to_numeric(selected["day_num"], errors="coerce")
            selected["Condition"] = str(condition)
            selected["Metric"] = option.column
            selected["RawValue"] = pd.to_numeric(selected["value"], errors="coerce")
            denominator = pd.to_numeric(selected["baseline_value"], errors="coerce")
            valid = denominator.notna() & denominator.ne(0)
            selected["Retention"] = float("nan")
            # H2 sensing retention follows the magnitude definition in the
            # previous JMP workbook. Other metrics use the signed ratio shown
            # by SensorLab's retention plots (100 - degradation_pct).
            if sheet == "JG_H2" and option.metric == "delta_Vth_H2":
                selected.loc[valid, "Retention"] = (
                    100 * selected.loc[valid, "RawValue"].abs()
                    / denominator.loc[valid].abs())
            else:
                selected.loc[valid, "Retention"] = (
                    100 * selected.loc[valid, "RawValue"] / denominator.loc[valid])
            parts.append(selected[_KEYS + ["Metric", "RawValue", "Retention"]])
    if not parts:
        raise ValueError("No rows were found for the selected metrics.")
    rows = pd.concat(parts, ignore_index=True)
    if rows["Day"].isna().any():
        raise ValueError("An exported source row has no numeric Day.")
    duplicate = rows.duplicated(_KEYS + ["Metric"], keep=False)
    if duplicate.any():
        example = rows.loc[duplicate, _KEYS + ["Metric"]].head(8)
        raise ValueError("Duplicate device/day/metric in analysis masters:\n"
                         + example.to_string(index=False))
    return rows


def make_export_tables(
    source: ExportSource, selected_metrics: list[MetricOption], *,
    stacking: str = "day",
) -> ExportTables:
    """One row per condition/device/day; stacking changes only row order."""
    if stacking not in {"day", "device"}:
        raise ValueError("stacking must be 'day' or 'device'.")
    valid = set(source.options)
    chosen = set(selected_metrics)
    if not chosen or not chosen <= valid:
        raise ValueError("Select at least one available metric.")
    rows = _selected_rows(source, chosen)
    ordered = [option for option in source.options if option in chosen]
    index = pd.MultiIndex.from_frame(rows[_KEYS].drop_duplicates())
    raw = rows.pivot(index=_KEYS, columns="Metric", values="RawValue")
    retention = rows.pivot(index=_KEYS, columns="Metric", values="Retention")
    raw = raw.reindex(index)
    retention = retention.reindex(index)
    columns: dict[str, pd.Series] = {}
    for option in ordered:
        columns[option.column] = (raw[option.column] if option.column in raw
                                  else pd.Series(float("nan"), index=index))
        name = option.column + "__Retention_pct"
        columns[name] = (retention[option.column] if option.column in retention
                         else pd.Series(float("nan"), index=index))
    data = pd.DataFrame(columns, index=index).reset_index()
    back_gates: set[str] = set()
    for option in ordered:
        explicit = dict(option.qualifiers).get("Vbg")
        suffix = re.search(r"_bg(-?\d+(?:\.\d+)?)$", option.metric)
        vbg = explicit or (suffix.group(1) if suffix else None)
        if vbg is None:
            break
        back_gates.add(vbg)
    else:
        if len(back_gates) == 1:
            data.insert(3, "Vbg", float(next(iter(back_gates))))
    # JG_H2 is the documented 100 ppm exposure. Only attach the concentration
    # as a row field when every selected metric belongs to the same n03/H2
    # experiment; STEPS metrics may refer to other concentrations.
    if any(option.sheet == "JG_H2" for option in ordered) and all(
            option.sheet in {"JG_H2", "JG_n03"} for option in ordered):
        data.insert(4 if "Vbg" in data else 3, "H2_ppm", 100)
    sort = (["Condition", "Day", "ChipID", "DeviceID"] if stacking == "day"
            else ["Condition", "ChipID", "DeviceID", "Day"])
    data = data.sort_values(sort, kind="stable").reset_index(drop=True)
    post = data.loc[data["Day"].gt(source.baseline_day)].reset_index(drop=True)

    summary_rows = rows.loc[rows["Day"].gt(source.baseline_day)].copy()
    if len(ordered) == 2:
        # Match the original two-metric JMP sheet: both checkpoint means use
        # exactly the same complete device pairs, rather than independent N.
        pair_columns = [option.column + "__Retention_pct" for option in ordered]
        paired = post.loc[post[pair_columns].notna().all(axis=1), _KEYS]
        summary_rows = summary_rows.merge(paired, on=_KEYS, how="inner",
                                          validate="many_to_one")
    summaries: list[dict[str, object]] = []
    for (condition, day, metric), group in summary_rows.groupby(
                ["Condition", "Day", "Metric"], sort=True):
        for measure, column in (("RawValue", "RawValue"),
                                ("Retention", "Retention_pct")):
            values = group[measure].dropna()
            summaries.append({"Condition": condition, "Day": day, "Metric": metric,
                              "Measure": column, "ValidN": int(len(values)),
                              "Mean": values.mean() if len(values) else float("nan"),
                              "SD": values.std(ddof=1) if len(values) > 1 else float("nan")})
    summary = pd.DataFrame(summaries, columns=["Condition", "Day", "Metric",
                                               "Measure", "ValidN", "Mean", "SD"])

    info = [
        ("All_Data / PostBaseline", "ChipID", "Chip identifier from ESN.", "identifier",
         "ESN + source ESN", "Nominal"),
        ("All_Data / PostBaseline", "DeviceID", "Stable transistor identifier.", "identifier",
         "ChipID + _T + source TNumber", "Nominal"),
        ("All_Data / PostBaseline", "Day", "Elapsed days; numeric storage.", "days",
         "day_num from source; model as a categorical checkpoint in JMP.", "Nominal"),
        ("All_Data / PostBaseline", "Condition", "Aging condition.", "text",
         "Condition from analysis master.", "Nominal"),
    ]
    if "Vbg" in data:
        info.append(("All_Data / PostBaseline", "Vbg", "Common back-gate voltage of all selected metrics.",
                     "V", "Parsed from source metric suffix or Vbg field.", "Continuous"))
    if "H2_ppm" in data:
        info.append(("All_Data / PostBaseline", "H2_ppm", "JG_H2 exposure concentration.",
                     "ppm", "100 ppm from the documented JG_H2 experiment; "
                     "not a numeric source column.", "Continuous"))
    for option in ordered:
        source_text = f"{option.sheet} analysis master; metric={option.metric}"
        if option.qualifiers:
            source_text += "; " + ", ".join(f"{k}={v}" for k, v in option.qualifiers)
        info.append(("All_Data / PostBaseline", option.column,
                     "Original per-device metric value, with source sign preserved.",
                     "source units", source_text + "; value (no rounding).", "Continuous"))
        formula = ("100 * abs(value(t)) / abs(value(0))"
                   if option.sheet == "JG_H2" and option.metric == "delta_Vth_H2"
                   else "100 * value(t) / value(0) (signed ratio)")
        info.append(("All_Data / PostBaseline", option.column + "__Retention_pct",
                     "Retention relative to the same device at the selected baseline.",
                     "% (numeric)", formula + "; blank if baseline missing/zero.",
                     "Continuous"))
    for name, desc, units, model in (
        ("Condition", "Aging condition.", "text", "Nominal"),
        ("Day", "Post-baseline checkpoint.", "days", "Nominal"),
        ("Metric", "Source metric with measurement context.", "text", "Nominal"),
        ("Measure", "RawValue or Retention_pct.", "text", "Nominal"),
        ("ValidN", "Number of nonmissing devices.", "devices", "Continuous"),
        ("Mean", "Arithmetic mean for this metric/measure/day.", "source units or %", "Continuous"),
        ("SD", "Sample standard deviation (n-1); blank for n<2.",
         "source units or percentage points", "Continuous"),
    ):
        info.append(("CheckpointSummary", name, desc, units,
                     "Calculated from the corresponding All_Data column.", model))
    column_info = pd.DataFrame(info, columns=["Sheet", "Column", "Description",
                                              "Units", "Formula_or_Source",
                                              "JMP_Modeling_Type"])
    notes = [
        f"Rows: All_Data={len(data)}, PostBaseline={len(post)}.",
        f"Devices={data['DeviceID'].nunique()}, conditions={data['Condition'].nunique()}, "
        f"days={sorted(data['Day'].unique().tolist())}.",
        f"Selected metrics={len(ordered)}; stacking={stacking}-major.",
        f"Baseline day={source.baseline_day:g}; PostBaseline excludes rows at/before it.",
        "Missing values remain blank. Zero/missing baseline denominators produce blank retention.",
        "JG_H2 delta_Vth_H2 retention uses magnitude; other retention uses signed value ratio.",
    ]
    if len(ordered) == 2:
        notes.append("CheckpointSummary uses the same complete retention pairs "
                     "for both selected metrics at each checkpoint.")
    for option in ordered:
        name = option.column
        notes.append(f"{name}: missing raw={int(data[name].isna().sum())}; "
                     f"missing retention={int(data[name + '__Retention_pct'].isna().sum())}.")
    return ExportTables(data, post, summary, column_info, notes)


def write_export(path: str | Path, tables: ExportTables) -> list[Path]:
    """Write one workbook plus two JMP-ready CSVs and a quality report."""
    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Choose an .xlsx output file.")
    target.parent.mkdir(parents=True, exist_ok=True)
    stem = target.with_suffix("")
    all_csv = Path(str(stem) + "_All_Data.csv")
    post_csv = Path(str(stem) + "_PostBaseline.csv")
    qc = Path(str(stem) + "_QC.txt")
    sheets = {"All_Data": tables.all_data, "PostBaseline": tables.post_baseline,
              "CheckpointSummary": tables.checkpoint_summary,
              "ColumnInfo": tables.column_info}
    # Build beside the destination so the final workbook appears only after
    # Excel writing succeeds. Existing files are touched only at the end.
    with tempfile.NamedTemporaryFile(prefix=".sensorlab_jmp_", suffix=".xlsx",
                                     dir=target.parent, delete=False) as tmp:
        temporary = Path(tmp.name)
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for sheet_name, frame in sheets.items():
                frame.to_excel(writer, sheet_name=sheet_name, index=False)
                ws = writer.sheets[sheet_name]
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for col_index, column in enumerate(frame.columns, start=1):
                    cell = ws.cell(1, col_index)
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center")
                    ws.column_dimensions[get_column_letter(col_index)].width = min(
                        max(len(str(column)) + 2, 13), 45)
                    if pd.api.types.is_numeric_dtype(frame[column]):
                        for row in ws.iter_rows(min_row=2, min_col=col_index,
                                                max_col=col_index):
                            row[0].number_format = "0.###############"
        os.replace(temporary, target)
    finally:
        if temporary.exists():
            temporary.unlink()
    tables.all_data.to_csv(all_csv, index=False, encoding="utf-8-sig",
                           float_format="%.17g")
    tables.post_baseline.to_csv(post_csv, index=False, encoding="utf-8-sig",
                               float_format="%.17g")
    qc.write_text("\n".join(tables.quality_notes) + "\n", encoding="utf-8")
    return [target, all_csv, post_csv, qc]
